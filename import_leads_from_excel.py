"""Import a provider Excel export through the running webhook (preview by default)."""
from __future__ import annotations

import argparse
from datetime import datetime
from pathlib import Path
from urllib.parse import parse_qs, urlsplit
from zoneinfo import ZoneInfo

import requests
from openpyxl import load_workbook


def _text(value) -> str:
    return "" if value is None else str(value).strip()


def build_payload(row: dict, source_tz: ZoneInfo) -> dict:
    uid = _text(row.get("id"))
    if not uid.isdigit():
        raise ValueError("id must contain digits")
    site = _text(row.get("Domain"))
    page = _text(row.get("Referer"))
    parsed = urlsplit(page)
    if not site or parsed.scheme not in {"http", "https"} or parsed.hostname != site.lower():
        raise ValueError("Referer must be an HTTP(S) page on Domain")
    phone = _text(row.get("Phone"))
    if not phone.lstrip("+").isdigit() or not 7 <= len(phone.lstrip("+")) <= 15:
        raise ValueError("Phone must contain 7 to 15 digits")
    created = row.get("Created")
    if not isinstance(created, datetime):
        created = datetime.strptime(_text(created), "%Y/%m/%d %H:%M:%S")
    if created.tzinfo is None:
        created = created.replace(tzinfo=source_tz)
    email = _text(row.get("Email"))
    payload = {
        "vid": uid,
        "site": site,
        "page": page,
        "time": str(int(created.timestamp())),
        "phones": [phone],
        "mails": [email] if email else [],
        "utm": {k: v[0] for k, v in parse_qs(parsed.query, keep_blank_values=True).items()},
        "comment": row.get("Comment"),
        "browser": _text(row.get("Browser")),
        "device": _text(row.get("Device")),
        "platform": _text(row.get("Platform")),
    }
    if row.get("Region"):
        payload["region"] = _text(row["Region"])
    return payload


def read_payloads(file: Path, source_tz: ZoneInfo) -> tuple[list[dict], int]:
    workbook = load_workbook(file, read_only=True, data_only=True)
    grouped: dict[str, dict] = {}
    count = 0
    try:
        if len(workbook.worksheets) != 1:
            raise ValueError("Expected one worksheet")
        rows = workbook.active.iter_rows(values_only=True)
        headers = [_text(v) for v in next(rows, ())]
        if len(set(headers)) != len(headers):
            raise ValueError("Duplicate column headers")
        required = {"id", "Domain", "Phone", "Created", "Referer"}
        if not required.issubset(headers):
            raise ValueError(f"Missing columns: {sorted(required - set(headers))}")
        for number, values in enumerate(rows, 2):
            if not any(v is not None for v in values):
                continue
            try:
                payload = build_payload(dict(zip(headers, values)), source_tz)
                uid = payload["vid"]
                existing = grouped.get(uid)
                if existing is None:
                    grouped[uid] = payload
                else:
                    for key in payload.keys() | existing.keys():
                        if key in {"phones", "mails"}:
                            continue
                        if payload.get(key) != existing.get(key):
                            raise ValueError(f"Conflicting {key} for repeated id={uid}")
                    for key in ("phones", "mails"):
                        existing[key].extend(v for v in payload[key] if v not in existing[key])
            except (ValueError, TypeError) as exc:
                raise ValueError(f"Excel row {number}: {exc}") from exc
            count += 1
    finally:
        workbook.close()
    return sorted(grouped.values(), key=lambda p: (int(p["time"]), p["vid"])), count


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--file", type=Path, required=True)
    parser.add_argument("--timezone", default="Europe/Moscow", help="Timezone of Created")
    parser.add_argument("--send", action="store_true", help="Save via webhook and queue emails")
    args = parser.parse_args()
    # Validate the entire file before the first request can send an email.
    payloads, rows = read_payloads(args.file, ZoneInfo(args.timezone))
    print(f"rows={rows} unique_leads={len(payloads)} phones={sum(len(p['phones']) for p in payloads)}")
    if not args.send:
        print("PREVIEW ONLY: no database writes or emails. Use --send on the VPS to import.")
        return

    from app.config import WEBHOOK_SECRET

    endpoint = f"http://127.0.0.1:8000/api/provider-test/{WEBHOOK_SECRET}"
    counts = {"new": 0, "duplicate": 0}
    with requests.Session() as session:
        # Loopback requests must not pass through an environment-configured proxy.
        session.trust_env = False
        for payload in payloads:
            uid = payload["vid"]
            try:
                response = session.post(endpoint, json=payload, timeout=30, allow_redirects=False)
            except requests.RequestException:
                raise SystemExit(f"Request failed for id={uid}; rerun safely with the same IDs") from None
            if response.status_code != 200:
                raise SystemExit(f"Webhook HTTP {response.status_code} for id={uid}; import stopped")
            try:
                result = response.json()
                state = result["lead_state"]
                if result.get("ok") is not True or state not in counts or not isinstance(result.get("lead_id"), int):
                    raise ValueError("Unexpected response")
            except (ValueError, KeyError, TypeError):
                raise SystemExit(f"Invalid webhook response for id={uid}; import stopped") from None
            counts[state] += 1
            print(f"id={uid} db_id={result['lead_id']} state={state}", flush=True)
    print(f"DONE new={counts['new']} duplicate={counts['duplicate']}")


if __name__ == "__main__":
    main()
